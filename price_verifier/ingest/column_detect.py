"""
column_detect.py — read a vendor's CSV/XLSX and work out which columns hold
the ASIN, the expected price and (optionally) the brand.

Why this exists: real vendor sheets do not have an `asin` /
`expected_price` / `brand` header row. They have "ASIN No.", "Amazon Link",
"SP", "Selling Price", "Brand Name", a "Price list – Sept 2026" title line
above the header, several sheets, MRP sitting next to SP, and so on. An
exact-header parser rejected those files outright. Instead we:

1. `load_table()`   — decode the file (CSV encoding/delimiter sniffing, or
                      openpyxl for XLSX with automatic data-sheet choice),
                      find the real header row, drop blank rows while
                      remembering each row's original row number.
2. `detect_columns()` — score every column for every field using BOTH the
                      header text and the cell contents, and return the best
                      mapping plus a ranked option list per field (for a
                      confirmation dropdown), confidences and vendor-facing
                      warnings.
3. `parse_rows()`   — turn a (possibly vendor-corrected) mapping into the
                      `ParseReport` the rest of the app consumes.

Content beats header text for the ASIN (an ASIN is unmistakable); for the
price, header text decides between numeric columns (MRP vs SP vs Qty all
parse as numbers). Everything here is pure Python + openpyxl, no network.
"""

from __future__ import annotations

import csv
import io
import math
import re
import statistics
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import load_workbook

from price_verifier.ingest.input_parser import (
    InputValidationError,
    ParsedRow,
    ParseReport,
)

__all__ = [
    "FIELDS",
    "RawTable",
    "ColumnOption",
    "DetectionResult",
    "load_table",
    "detect_columns",
    "extract_asin",
    "parse_price",
    "parse_rows",
    "raise_if_unmapped",
    "display_value",
    # re-exported for convenience
    "InputValidationError",
    "ParsedRow",
    "ParseReport",
]

FIELDS = ("asin", "expected_price", "brand")  # brand is OPTIONAL

# ── Tunables ────────────────────────────────────────────────────────────────
HEADER_SCAN_ROWS = 15          # non-empty rows examined when looking for the header
HEADER_LOOKAHEAD_ROWS = 10     # rows under a header candidate checked for ASIN data
DETECTION_SAMPLE_ROWS = 1000   # rows used for column scoring (keeps 10k-row files fast)
SHEET_SCAN_ROWS = 200          # rows per sheet read when auto-picking the data sheet
PREVIEW_ROWS = 5
MAX_SAMPLES = 3

ASIN_MIN_CONTENT = 0.5         # share of non-empty cells that must be ASINs/links to auto-map
PRICE_MIN_CONTENT = 0.5        # share of non-empty cells that must parse as a price
STRONG_HEADER = 0.8            # header score at/above which the name alone is convincing
AMBIGUITY_MARGIN = 0.1         # runner-up within this score of the winner => "please confirm"

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .xls (and other OLE2) files
_CSV_DELIMITERS = (",", ";", "\t", "|")


# ── Public data types ───────────────────────────────────────────────────────

@dataclass
class RawTable:
    sheet_names: list[str]      # [] for CSV
    sheet_name: Optional[str]
    header_row: int             # 1-based row of the header line; 0 = file has no header row
    headers: list[str]          # display names; "Column A", "Column B"... for blank/no header
    rows: list[list]            # data rows after the header, raw cell values, fully-blank rows removed
    # Original 1-based file row number for each entry in `rows` (blank rows
    # and title lines removed, so this is not simply header_row + 1 + i).
    row_numbers: list[int] = field(default_factory=list)
    # Non-empty lines above the header (or above the data, if headerless)
    # that were skipped as titles/banners, as display text.
    title_rows: list[str] = field(default_factory=list)
    sheet_auto_picked: bool = False

    def row_number(self, i: int) -> int:
        if i < len(self.row_numbers):
            return self.row_numbers[i]
        return self.header_row + 1 + i


@dataclass
class ColumnOption:
    index: int
    header: str
    score: float
    reason: str
    samples: list[str]  # up to 3 sample display values


@dataclass
class DetectionResult:
    table: RawTable
    mapping: dict[str, Optional[int]]        # field -> column index (None = not detected / not mapped)
    confidence: dict[str, float]             # 0..1 per field
    options: dict[str, list[ColumnOption]]   # per field: ALL columns ranked best-first
    warnings: list[str]                      # human-readable, vendor-friendly
    preview: list[list[str]]                 # first 5 data rows as display strings


# ── Cell helpers ────────────────────────────────────────────────────────────

def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def display_value(v) -> str:
    """How a raw cell value should be shown to the vendor."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if math.isfinite(v) and v.is_integer():
            return str(int(v))
        return format(v, ".10g")
    if isinstance(v, datetime):
        if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0):
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _column_letter(i: int) -> str:
    s = ""
    n = i + 1
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


# ── ASIN extraction ─────────────────────────────────────────────────────────

_ASIN_B_RE = re.compile(r"^B[0-9A-Z]{9}$")
_ASIN_ISBN_RE = re.compile(r"^\d{9}[\dX]$")
_URL_HINT_RE = re.compile(r"amazon\.|amzn\.|/dp/|/gp/|[?&]asin=", re.I)
_URL_ASIN_RE = re.compile(
    r"(?:/dp/|/gp/product/|/gp/aw/d/|/product/|[?&]asin=)([A-Za-z0-9]{10})(?![A-Za-z0-9])",
    re.I,
)
_FLOATY_DIGITS_RE = re.compile(r"^(\d{9,10})\.0+$")
_STRIP_CHARS = " \t\r\n\"'` ​﻿“”‘’"
_TRAILING_PUNCT = ".,;:!?)]}>"


def _is_valid_asin(s: str) -> bool:
    return bool(_ASIN_B_RE.match(s) or _ASIN_ISBN_RE.match(s))


def _asin_from_number(n: int) -> Optional[str]:
    # Excel stores an ISBN-style ASIN typed without text formatting as a
    # number, dropping a leading zero ("0123456789" -> 123456789).
    digits = str(n)
    if len(digits) not in (9, 10):
        return None
    s = digits.zfill(10)
    return s if _ASIN_ISBN_RE.match(s) else None


def extract_asin(value) -> Optional[str]:
    """Return the ASIN in `value` (a plain ASIN or an Amazon product URL),
    upper-cased, or None if there isn't one."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return _asin_from_number(value) if value > 0 else None
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer() or value <= 0:
            return None
        return _asin_from_number(int(value))
    s = str(value).strip(_STRIP_CHARS)
    if not s:
        return None
    if _URL_HINT_RE.search(s):
        m = _URL_ASIN_RE.search(s)
        if not m:
            return None
        cand = m.group(1).upper()
        return cand if _is_valid_asin(cand) else None
    m = _FLOATY_DIGITS_RE.match(s)
    if m:
        return _asin_from_number(int(m.group(1)))
    s = s.rstrip(_TRAILING_PUNCT).strip(_STRIP_CHARS).upper()
    return s if _is_valid_asin(s) else None


def _asin_error_reason(value) -> str:
    """Vendor-friendly explanation for why extract_asin(value) was None."""
    if _is_blank(value):
        return "ASIN is blank"
    s = display_value(value)
    if _URL_HINT_RE.search(s) or s.lower().startswith(("http", "www.")):
        return "Couldn't find an ASIN in this link — use the full Amazon product link (with /dp/…) or the ASIN itself"
    s = s.strip(_STRIP_CHARS).rstrip(_TRAILING_PUNCT).strip(_STRIP_CHARS).upper()
    if len(s) != 10:
        return f"ASIN must be 10 characters, got {len(s)}"
    if not s.startswith("B") and not s[:9].isdigit():
        return "ASIN must start with 'B' (or be a 10-digit ISBN)"
    return "ASIN contains invalid characters"


# ── Price parsing ───────────────────────────────────────────────────────────

_PRICE_PREFIX_RE = re.compile(r"^(?:₹|rs\.?|inr|\$)\s*", re.I)
_PRICE_SUFFIX_RE = re.compile(r"\s*(?:/-|/=|₹|rs\.?|inr|only|\.)$", re.I)
_PLAIN_NUMBER_RE = re.compile(r"^\d+(?:\.\d+)?$")
# Western (1,499 / 1,234,567) and Indian lakh (1,49,999) grouping; the last
# group is always 3 digits so "1,2" (a European decimal) is rejected.
_GROUPED_NUMBER_RE = re.compile(r"^\d{1,3}(?:,\d{2,3})*,\d{3}(?:\.\d+)?$")


def parse_price(value) -> Optional[float]:
    """Parse a price cell ("₹1,499.00", "Rs 1,499/-", "INR 1499", 1499.5,
    "1,49,999") to a positive float. None for text, blank, zero, negative."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        f = float(value)
        return f if math.isfinite(f) and f > 0 else None
    if isinstance(value, (datetime, date)):
        return None
    s = str(value).replace(" ", " ").strip()
    if not s:
        return None
    for _ in range(3):
        new = _PRICE_PREFIX_RE.sub("", s)
        if new == s:
            break
        s = new
    for _ in range(4):
        new = _PRICE_SUFFIX_RE.sub("", s)
        if new == s:
            break
        s = new
    s = s.replace(" ", "")
    if _PLAIN_NUMBER_RE.match(s):
        f = float(s)
    elif _GROUPED_NUMBER_RE.match(s):
        f = float(s.replace(",", ""))
    else:
        return None
    return f if math.isfinite(f) and f > 0 else None


_ZERO_PRICE_RE = re.compile(r"(?:₹|rs\.?|inr)?\s*0+(?:\.0+)?\s*(?:/-)?", re.I)


def _price_error_reason(value) -> str:
    if _is_blank(value):
        return "Expected price is blank"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "Expected price must be more than 0"
    s = display_value(value)
    if s.startswith("-") and parse_price(s.lstrip("- ")) is not None:
        return "Expected price must be more than 0"
    if _ZERO_PRICE_RE.fullmatch(s):
        return "Expected price must be more than 0"
    return f"Expected price '{s[:40]}' is not a number"


# ── Header-name vocabulary ──────────────────────────────────────────────────

def _normalize_header(h) -> str:
    """'Expected_Price' -> 'expected price', 'S.P.' -> 'sp', 'ASIN No.' -> 'asin no'."""
    s = re.sub(r"([a-z])([A-Z])", r"\1 \2", display_value(h))
    tokens = re.sub(r"[^a-z0-9]+", " ", s.lower()).split()
    merged: list[str] = []
    run = ""
    for t in tokens:  # collapse dotted abbreviations: s p -> sp, m r p -> mrp
        if len(t) == 1 and t.isalpha():
            run += t
            continue
        if run:
            merged.append(run)
            run = ""
        merged.append(t)
    if run:
        merged.append(run)
    return " ".join(merged)


def _contains_phrase(norm: str, phrases: Iterable[str]) -> Optional[str]:
    padded = f" {norm} "
    for p in phrases:
        if f" {p} " in padded:
            return p
    return None


_ASIN_STRONG = ("asin", "asins", "asin no", "asin number", "asin code", "asin id", "amazon asin")
_ASIN_LINKISH = ("amazon link", "product link", "amazon url", "product url", "asin link",
                 "url", "link", "links", "amazon", "hyperlink")

_PRICE_STRONG = ("expected price", "expected", "agreed price", "agreed", "selling price", "sp",
                 "offer price", "our price", "target price", "net price", "deal price", "rate",
                 "price to check", "check price", "sale price", "sell price", "selling",
                 "final price", "special price", "expected sp")
_PRICE_WEAK = ("price", "amount", "amt", "value", "inr", "rs")
# Price-like but NOT the selling price — usable only as a last resort.
_PRICE_NEG_PRICELIKE = ("mrp", "max retail", "maximum retail", "retail price", "list price",
                        "cost", "cp", "purchase", "landing", "nlc", "dp", "dealer price")
# Numeric but never a price.
_PRICE_NEG_OTHER = ("margin", "discount", "disc", "off", "qty", "quantity", "stock", "gst", "tax",
                    "hsn", "weight", "sku", "count", "pincode", "pin", "zip", "postal", "sr",
                    "sno", "s no", "sl no", "serial", "rank", "ranking", "rating", "ratings",
                    "review", "reviews", "mobile", "phone", "ean", "upc", "barcode", "moq",
                    "units", "unit", "pack", "id", "code", "year", "days", "percent", "percentage")

_BRAND_STRONG = ("brand", "brands", "brand name", "make", "manufacturer", "mfr", "company",
                 "company name", "label", "oem")

_PINCODE_HEADERS = ("pincode", "pin code", "pin", "zip", "zipcode", "zip code", "postal code", "postcode")

_HEADER_VOCAB = frozenset(
    t
    for group in (_ASIN_STRONG, _ASIN_LINKISH, _PRICE_STRONG, _PRICE_WEAK, _PRICE_NEG_PRICELIKE,
                  _PRICE_NEG_OTHER, _BRAND_STRONG, _PINCODE_HEADERS,
                  ("name", "title", "product", "item", "description", "category", "model",
                   "seller", "no", "number", "remarks", "status", "sub category", "colour",
                   "color", "size", "type", "group", "details"))
    for phrase in group
    for t in phrase.split()
    if len(t) > 1
)


def _asin_header_score(norm: str) -> float:
    if norm in _ASIN_STRONG:
        return 1.0
    if _contains_phrase(norm, _ASIN_STRONG):
        return 0.9
    if norm in _ASIN_LINKISH:
        return 0.8
    if _contains_phrase(norm, _ASIN_LINKISH):
        return 0.7
    return 0.0


def _price_header_score(norm: str, raw: str) -> float:
    """1.0 exact strong / 0.9 contains strong / 0.5 weak / 0.4 contains weak /
    0 unknown / -0.5 price-like negative (MRP, cost) / -1 non-price numeric."""
    if not norm:
        return 0.0
    neg_other = _contains_phrase(norm, _PRICE_NEG_OTHER) or ("%" in raw)
    neg_price = _contains_phrase(norm, _PRICE_NEG_PRICELIKE)
    if norm in _PRICE_STRONG:
        return 1.0
    if _contains_phrase(norm, _PRICE_STRONG):
        return 0.3 if (neg_other or neg_price) else 0.9
    if neg_other:
        return -1.0
    if neg_price:
        return -0.5
    if norm in _PRICE_WEAK:
        return 0.5
    if _contains_phrase(norm, _PRICE_WEAK):
        return 0.4
    return 0.0


def _brand_header_score(norm: str) -> float:
    if norm in _BRAND_STRONG:
        return 1.0
    if _contains_phrase(norm, _BRAND_STRONG):
        return 0.8
    return 0.0


def _vocab_hits(cells: list) -> int:
    hits = 0
    for c in cells:
        if isinstance(c, str) and c.strip():
            if any(t in _HEADER_VOCAB for t in _normalize_header(c).split()):
                hits += 1
    return hits


def _is_texty(v) -> bool:
    """Non-empty text that is neither a number, a price nor an ASIN."""
    return (isinstance(v, str) and bool(v.strip())
            and parse_price(v) is None and extract_asin(v) is None)


# ── Loading ─────────────────────────────────────────────────────────────────

def load_table(filename: str, data: bytes, sheet_name: Optional[str] = None) -> RawTable:
    """Read a .csv/.tsv/.xlsx/.xlsm upload into a RawTable (header row found,
    blank rows dropped). Raises InputValidationError with a vendor-friendly
    message for unsupported, unreadable or empty files."""
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xls" or (suffix in (".xlsx", ".xlsm") and data[:8] == _OLE_MAGIC):
        raise InputValidationError(
            "This is an old-style Excel file (.xls), which can't be read. Open it in Excel, "
            "choose File → Save As, pick 'Excel Workbook (.xlsx)' or 'CSV', and upload that file."
        )
    if suffix not in (".csv", ".tsv", ".xlsx", ".xlsm"):
        shown = suffix or "(no extension)"
        raise InputValidationError(
            f"Unsupported file type '{shown}'. Please upload an Excel file (.xlsx) or a .csv file."
        )
    if not data:
        raise InputValidationError("The file is empty — there's nothing in it to check.")

    if suffix in (".csv", ".tsv"):
        rows = _read_csv(data)
        return _build_table(rows, sheet_names=[], sheet_name=None, auto_picked=False)
    return _load_xlsx(data, sheet_name)


def _decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):  # Excel "Unicode Text" export
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pick_delimiter(text: str) -> str:
    lines = [ln for ln in text.splitlines() if ln.strip()][:50]
    if not lines:
        return ","
    sample = "\n".join(lines)
    try:
        sniffed = csv.Sniffer().sniff(sample, delimiters="".join(_CSV_DELIMITERS)).delimiter
    except csv.Error:
        sniffed = None

    def consistency(delim: str) -> tuple[int, int]:
        counts = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=delim)]
        multi = [c for c in counts if c > 1]
        if not multi:
            return (0, 0)
        mode = statistics.mode(multi)
        return (sum(1 for c in counts if c == mode), mode)

    scored = {d: consistency(d) for d in _CSV_DELIMITERS}
    best = max(scored.values())
    if best[0] == 0:
        return ","
    winners = [d for d in _CSV_DELIMITERS if scored[d] == best]
    if sniffed in winners:
        return sniffed
    return winners[0]


def _read_csv(data: bytes) -> list[tuple[int, list]]:
    text = _decode_text(data).replace("\x00", "")
    if text.startswith("﻿"):
        text = text[1:]
    delim = _pick_delimiter(text)
    try:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
        return [(i, row) for i, row in enumerate(reader, start=1)]
    except csv.Error as e:
        raise InputValidationError(
            "We couldn't read this CSV file — it may be damaged. Try opening it in Excel and "
            "saving it again as CSV or .xlsx."
        ) from e


def _read_sheet(ws, max_nonempty: Optional[int] = None) -> list[tuple[int, list]]:
    if hasattr(ws, "reset_dimensions"):
        # Some exporters write a wrong <dimension> tag; read_only mode would
        # then silently truncate the sheet. Resetting makes openpyxl read
        # every row actually present.
        ws.reset_dimensions()
    out: list[tuple[int, list]] = []
    nonempty = 0
    for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        cells = list(row)
        out.append((i, cells))
        if any(not _is_blank(c) for c in cells):
            nonempty += 1
            if max_nonempty is not None and nonempty >= max_nonempty:
                break
    return out


def _count_asin_rows(rows: list[tuple[int, list]]) -> int:
    return sum(1 for _, cells in rows if any(extract_asin(c) for c in cells))


def _load_xlsx(data: bytes, sheet_name: Optional[str]) -> RawTable:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # BadZipFile, InvalidFileException, KeyError, ...
        raise InputValidationError(
            "We couldn't open this Excel file. It may be damaged or password-protected — open it "
            "in Excel, save a fresh copy as .xlsx (without a password), and upload that."
        ) from e
    try:
        sheets = list(wb.worksheets)
        names = [ws.title for ws in sheets]
        if not sheets:
            raise InputValidationError("This Excel file has no sheets with data in it.")
        auto = False
        if sheet_name is not None:
            chosen = next((ws for ws in sheets if ws.title == sheet_name), None)
            if chosen is None:
                wanted = sheet_name.strip().lower()
                chosen = next((ws for ws in sheets if ws.title.strip().lower() == wanted), None)
            if chosen is None:
                raise InputValidationError(
                    f"There's no sheet called '{sheet_name}' in this file. "
                    f"Sheets found: {', '.join(names)}."
                )
        elif len(sheets) == 1:
            chosen = sheets[0]
        else:
            auto = True
            chosen, best_score, first_nonempty = None, 0, None
            for ws in sheets:
                sample = _read_sheet(ws, max_nonempty=SHEET_SCAN_ROWS)
                has_data = any(any(not _is_blank(c) for c in cells) for _, cells in sample)
                if has_data and first_nonempty is None:
                    first_nonempty = ws
                score = _count_asin_rows(sample)
                if score > best_score:
                    chosen, best_score = ws, score
            chosen = chosen or first_nonempty or sheets[0]
        rows = _read_sheet(chosen)
        return _build_table(rows, sheet_names=names, sheet_name=chosen.title, auto_picked=auto)
    finally:
        wb.close()


def _best_asin_column(rows: list[list], ncols: int) -> tuple[int, float]:
    if not rows:
        return 0, 0.0
    best_col, best_frac = 0, 0.0
    for c in range(ncols):
        hits = sum(1 for r in rows if c < len(r) and extract_asin(r[c]) is not None)
        frac = hits / len(rows)
        if frac > best_frac:
            best_col, best_frac = c, frac
    return best_col, best_frac


_NEAR_ASIN_RE = re.compile(r"B[0-9A-Z]{5,13}")


def _looks_like_bad_asin(v) -> bool:
    """A mistyped ASIN ("B09W9F") in a headerless file's first row must not
    be taken for a column heading."""
    if not isinstance(v, str):
        return False
    s = v.strip(_STRIP_CHARS).upper()
    return bool(_NEAR_ASIN_RE.fullmatch(s)) and any(ch.isdigit() for ch in s)


def _row_profile(cells: list) -> tuple[int, float, int]:
    """(non-empty count, text fraction, header-vocabulary hits)."""
    nonempty = [c for c in cells if not _is_blank(c)]
    if not nonempty:
        return 0, 0.0, 0
    text = sum(1 for c in nonempty if _is_texty(c))
    return len(nonempty), text / len(nonempty), _vocab_hits(nonempty)


def _detect_header(rows: list[list], ncols: int) -> tuple[Optional[int], int]:
    """Return (header index or None, index of first data row) within `rows`
    (the non-empty rows of the sheet)."""
    scan = rows[:HEADER_SCAN_ROWS]
    profiles = [_row_profile(r) for r in scan]

    # 1. A text-heavy row directly above rows that contain ASINs.
    best: Optional[tuple[float, int]] = None
    for i, cells in enumerate(scan):
        nonempty, text_frac, hits = profiles[i]
        if nonempty < 2 or text_frac < 0.6:
            continue
        following = rows[i + 1 : i + 1 + HEADER_LOOKAHEAD_ROWS]
        col, frac = _best_asin_column(following, ncols)
        if frac < 0.5:
            continue
        hdr_cell = cells[col] if col < len(cells) else None
        if extract_asin(hdr_cell) is not None or _looks_like_bad_asin(hdr_cell):
            continue  # a (malformed) data row, not a heading
        widest = max(sum(1 for c in r if not _is_blank(c)) for r in following)
        coverage = min(1.0, nonempty / max(1, widest))
        next_is_data = col < len(following[0]) and extract_asin(following[0][col]) is not None
        score = (min(hits, 3) + text_frac + coverage
                 + (0.5 if next_is_data else 0.0)
                 - (0.5 if _is_blank(hdr_cell) else 0.0))
        if best is None or score > best[0]:
            best = (score, i)
    if best is not None:
        return best[1], best[1] + 1

    has_asin_rows = [i for i, cells in enumerate(scan) if any(extract_asin(c) for c in cells)]

    # 2. No usable ASIN data under any candidate: a row naming known columns.
    for i, (nonempty, text_frac, hits) in enumerate(profiles):
        if has_asin_rows and i >= has_asin_rows[0]:
            break
        if nonempty >= 2 and text_frac >= 0.6 and hits >= 1:
            return i, i + 1

    # 3. Headerless: data starts at the first row carrying an ASIN, but a
    #    full-width row before it is data with a mistyped ASIN, not a title —
    #    keep it so the vendor gets a row-level error for it.
    if has_asin_rows:
        first = has_asin_rows[0]
        data_width = profiles[first][0]
        start = first
        while start > 0 and profiles[start - 1][0] >= max(2, math.ceil(0.75 * data_width)):
            start -= 1
        return None, start

    # 4. Classic layout with unfamiliar names: an all-text first row.
    nonempty, text_frac, _ = profiles[0]
    if nonempty >= 2 and text_frac == 1.0 and len(rows) > 1:
        return 0, 1
    return None, 0


def _build_table(numbered_rows: list[tuple[int, list]], sheet_names: list[str],
                 sheet_name: Optional[str], auto_picked: bool) -> RawTable:
    nonempty = [(n, cells) for n, cells in numbered_rows if any(not _is_blank(c) for c in cells)]
    where = f"The sheet '{sheet_name}'" if sheet_name else "The file"
    if not nonempty:
        raise InputValidationError(f"{where} is empty — there are no rows to check.")

    ncols = 0
    for _, cells in nonempty:
        for j in range(len(cells) - 1, ncols - 1, -1):
            if not _is_blank(cells[j]):
                ncols = j + 1
                break
    rows = [(cells + [None] * (ncols - len(cells)))[:ncols] for _, cells in nonempty]
    numbers = [n for n, _ in nonempty]

    hdr_idx, data_start = _detect_header(rows, ncols)
    title_rows = [
        " ".join(display_value(c) for c in rows[i] if not _is_blank(c))
        for i in range(0, hdr_idx if hdr_idx is not None else data_start)
    ]
    if hdr_idx is not None:
        headers = [display_value(h) or f"Column {_column_letter(j)}" for j, h in enumerate(rows[hdr_idx])]
        header_row = numbers[hdr_idx]
    else:
        headers = [f"Column {_column_letter(j)}" for j in range(ncols)]
        header_row = 0

    data_rows = rows[data_start:]
    if not data_rows:
        raise InputValidationError(
            f"{where} has a header row but no data rows under it — add your ASINs and prices below the headings."
        )
    return RawTable(
        sheet_names=sheet_names,
        sheet_name=sheet_name,
        header_row=header_row,
        headers=headers,
        rows=data_rows,
        row_numbers=numbers[data_start:],
        title_rows=title_rows,
        sheet_auto_picked=auto_picked,
    )


# ── Column statistics & scoring ─────────────────────────────────────────────

@dataclass
class _ColStats:
    index: int
    header: str
    norm: str
    total: int
    non_empty: int = 0
    asin_hits: int = 0
    asin_b_hits: int = 0
    price_hits: int = 0
    text_hits: int = 0
    distinct: int = 0
    median_len: float = 0.0
    sequential: bool = False
    pincode_like: bool = False
    samples: list[str] = field(default_factory=list)

    @property
    def fill(self) -> float:
        return self.non_empty / self.total if self.total else 0.0

    def frac(self, hits: int) -> float:
        return hits / self.non_empty if self.non_empty else 0.0


def _column_stats(table: RawTable, sample: list[list]) -> list[_ColStats]:
    out = []
    for j, header in enumerate(table.headers):
        norm = _normalize_header(header) if table.header_row else ""
        st = _ColStats(index=j, header=header, norm=norm, total=len(sample))
        values, lengths, numbers = set(), [], []
        for r in sample:
            v = r[j] if j < len(r) else None
            if _is_blank(v):
                continue
            st.non_empty += 1
            shown = display_value(v)
            if len(st.samples) < MAX_SAMPLES:
                st.samples.append(shown)
            asin = extract_asin(v)
            if asin is not None:
                st.asin_hits += 1
                if asin.startswith("B"):
                    st.asin_b_hits += 1
            price = parse_price(v)
            if price is not None:
                st.price_hits += 1
                numbers.append(price)
            elif asin is None and isinstance(v, str):
                st.text_hits += 1
            values.add(shown.lower())
            lengths.append(len(shown))
        st.distinct = len(values)
        st.median_len = statistics.median(lengths) if lengths else 0.0
        if len(numbers) >= 3:
            steps = sum(1 for a, b in zip(numbers, numbers[1:]) if b - a == 1)
            st.sequential = steps >= 0.8 * (len(numbers) - 1)
            st.pincode_like = all(n.is_integer() and 100000 <= n <= 999999 for n in numbers)
        out.append(st)
    return out


def _pct(x: float) -> str:
    return f"{round(x * 100)}%"


def _score_asin(st: _ColStats) -> tuple[float, float, str]:
    """(score, content fraction, reason)."""
    h = _asin_header_score(st.norm)
    a = st.frac(st.asin_hits)
    score = 0.75 * a + 0.25 * h
    if st.asin_hits and not st.asin_b_hits and h < 0.7:
        score *= 0.8  # all digit-only (ISBN-style) — could be phone numbers etc.
    parts = []
    if st.non_empty:
        parts.append(f"{_pct(a)} of values are ASINs or Amazon links")
    else:
        parts.append("column is empty")
    if h >= 0.7:
        parts.append(f"heading '{st.header}' suggests ASINs")
    return round(min(score, 1.0), 3), a, "; ".join(parts)


def _score_price(st: _ColStats) -> tuple[float, float, float, str]:
    """(score, content fraction, header score, reason)."""
    h = _price_header_score(st.norm, st.header if st.norm else "")
    f = st.frac(st.price_hits)
    content = f * (0.5 + 0.5 * st.fill)
    notes = []
    if st.sequential:
        content *= 0.3
        notes.append("looks like a serial number")
    if st.pincode_like and h <= 0:
        content *= 0.3
        notes.append("looks like pincodes")
    if h >= 0:
        score = 0.5 * content + 0.5 * h
    elif h == -0.5:
        score = 0.25 * content
    else:
        score = 0.05 * content
    parts = [f"{_pct(f)} of values are prices" if st.non_empty else "column is empty"]
    if h >= STRONG_HEADER:
        parts.append(f"heading '{st.header}' means the selling price")
    elif h > 0:
        parts.append(f"heading '{st.header}' is a general price/amount")
    elif h == -0.5:
        parts.append(f"heading '{st.header}' looks like MRP/cost, not the selling price")
    elif h < 0:
        parts.append(f"heading '{st.header}' isn't a price")
    parts.extend(notes)
    return round(min(score, 1.0), 3), f, h, "; ".join(parts)


def _score_brand(st: _ColStats) -> tuple[float, float, float, str]:
    """(score, content score, header score, reason)."""
    h = _brand_header_score(st.norm)
    t = st.frac(st.text_hits)
    if st.non_empty < 10:
        card = 0.5
    else:
        ratio = st.distinct / st.non_empty
        card = max(0.0, min(1.0, (0.9 - ratio) / 0.6))
    length = max(0.0, min(1.0, (50 - st.median_len) / 30))
    content = t * (0.5 * card + 0.5 * length)
    score = 0.6 * h + 0.4 * content
    parts = []
    if st.non_empty:
        parts.append(f"{_pct(t)} of values are text")
        if st.non_empty >= 10:
            parts.append(f"{st.distinct} different values")
    else:
        parts.append("column is empty")
    if h > 0:
        parts.append(f"heading '{st.header}' suggests brand")
    return round(min(score, 1.0), 3), content, h, "; ".join(parts)


def _ranked(options: list[ColumnOption]) -> list[ColumnOption]:
    return sorted(options, key=lambda o: (-o.score, o.index))


# ── Detection ───────────────────────────────────────────────────────────────

def detect_columns(table: RawTable) -> DetectionResult:
    """Work out which column holds each of FIELDS. Never raises: fields that
    can't be found map to None (see raise_if_unmapped for the strict check)."""
    sample = table.rows[:DETECTION_SAMPLE_ROWS]
    stats = _column_stats(table, sample)
    warnings: list[str] = []
    mapping: dict[str, Optional[int]] = {f: None for f in FIELDS}
    confidence: dict[str, float] = {f: 0.0 for f in FIELDS}
    options: dict[str, list[ColumnOption]] = {}

    if table.sheet_auto_picked and len(table.sheet_names) > 1:
        warnings.append(
            f"Your file has {len(table.sheet_names)} sheets — we used '{table.sheet_name}'. "
            f"Choose a different sheet if that's not the right one."
        )
    if table.title_rows:
        example = table.title_rows[0][:60]
        n = len(table.title_rows)
        warnings.append(
            f"Skipped {n} line{'s' if n != 1 else ''} above the "
            f"{'column headings' if table.header_row else 'data'} (e.g. '{example}')."
        )
    if table.header_row == 0:
        warnings.append(
            "Your file doesn't seem to have a row of column headings, so we worked out the "
            "columns from their contents. Please check them."
        )

    # ── ASIN ──
    asin_opts, asin_eligible = [], []
    for st in stats:
        score, a, reason = _score_asin(st)
        asin_opts.append(ColumnOption(st.index, st.header, score, reason, st.samples))
        if a >= ASIN_MIN_CONTENT:
            asin_eligible.append((score, st))
    options["asin"] = _ranked(asin_opts)
    if asin_eligible:
        score, st = max(asin_eligible, key=lambda x: (x[0], -x[1].index))
        mapping["asin"], confidence["asin"] = st.index, score
    else:
        fallback = [st for st in stats if _asin_header_score(st.norm) == 1.0]
        if fallback:
            st = fallback[0]
            mapping["asin"], confidence["asin"] = st.index, 0.3
            if st.non_empty:
                warnings.append(
                    f"Only {_pct(st.frac(st.asin_hits))} of the values in '{st.header}' look like "
                    f"valid ASINs (e.g. B09W9FND7M). Please check this is the right column."
                )
    asin_col = mapping["asin"]

    # ── Expected price ──
    price_opts, price_eligible = [], []
    for st in stats:
        if st.index == asin_col:
            price_opts.append(ColumnOption(st.index, st.header, 0.0, "already used for the ASIN", st.samples))
            continue
        score, f, h, reason = _score_price(st)
        price_opts.append(ColumnOption(st.index, st.header, score, reason, st.samples))
        if f >= PRICE_MIN_CONTENT:
            price_eligible.append((score, h, st))
    options["expected_price"] = _ranked(price_opts)
    if price_eligible:
        price_eligible.sort(key=lambda x: (-x[0], x[2].index))
        score, h, st = price_eligible[0]
        mapping["expected_price"] = st.index
        conf = score
        runner = price_eligible[1] if len(price_eligible) > 1 else None
        if runner and runner[1] >= STRONG_HEADER and h >= STRONG_HEADER and score - runner[0] <= AMBIGUITY_MARGIN:
            conf *= 0.8
            warnings.append(
                f"Both '{st.header}' and '{runner[2].header}' look like the selling price — we used "
                f"'{st.header}'. Please confirm."
            )
        if h == -0.5:
            warnings.append(
                f"We couldn't find a selling-price column, so we used '{st.header}'. That usually "
                f"holds the MRP or cost, which is often higher than the price on Amazon — please "
                f"make sure it's the price you expect Amazon to show."
            )
        elif h < 0:
            warnings.append(
                f"We couldn't find a price column, so we used '{st.header}' because it holds "
                f"numbers. Please check it really is the price you expect Amazon to show."
            )
        elif h < STRONG_HEADER:
            label = f"'{st.header}'"
            warnings.append(
                f"We're using {label} as the expected price, but its heading doesn't say "
                f"'selling price' — please check it's the price you expect Amazon to show (not MRP)."
            )
        confidence["expected_price"] = round(min(conf, 1.0), 3)
    else:
        fallback = [st for st in stats
                    if st.index != asin_col and _price_header_score(st.norm, st.header) == 1.0]
        if fallback:
            st = fallback[0]
            mapping["expected_price"], confidence["expected_price"] = st.index, 0.3
            if st.non_empty:
                warnings.append(
                    f"Most values in '{st.header}' aren't numbers — rows without a usable price "
                    f"will be skipped. Please check this is the right column."
                )
    price_col = mapping["expected_price"]

    # ── Brand (optional) ──
    brand_opts, brand_candidates = [], []
    for st in stats:
        if st.index in (asin_col, price_col):
            used = "ASIN" if st.index == asin_col else "expected price"
            brand_opts.append(ColumnOption(st.index, st.header, 0.0, f"already used for the {used}", st.samples))
            continue
        score, content, h, reason = _score_brand(st)
        brand_opts.append(ColumnOption(st.index, st.header, score, reason, st.samples))
        brand_candidates.append((score, content, h, st))
    options["brand"] = _ranked(brand_opts)

    by_header = [c for c in brand_candidates
                 if c[2] >= STRONG_HEADER and (c[3].non_empty == 0 or c[3].frac(c[3].text_hits) >= 0.5)]
    if by_header:
        score, _, _, st = max(by_header, key=lambda c: (c[0], -c[3].index))
        mapping["brand"], confidence["brand"] = st.index, score
    elif table.header_row == 0:
        # No headings at all: accept a clearly brand-shaped text column
        # (mostly short text, few distinct values), but only if unambiguous.
        shaped = sorted(
            (c for c in brand_candidates
             if c[3].non_empty >= 10 and c[3].frac(c[3].text_hits) >= 0.8
             and c[3].distinct / c[3].non_empty <= 0.5 and c[3].median_len <= 30 and c[1] >= 0.6),
            key=lambda c: -c[1],
        )
        if shaped and (len(shaped) == 1 or shaped[0][1] - shaped[1][1] >= 0.15):
            _, content, _, st = shaped[0]
            mapping["brand"], confidence["brand"] = st.index, round(0.6 * content, 3)
            warnings.append(f"We guessed that '{st.header}' holds the brand — please check.")
    if mapping["brand"] is None:
        warnings.append(
            "No brand column was found. Brand is optional — if your file has one, select it."
        )

    # ── Duplicate ASINs (full file, not just the sample) ──
    if asin_col is not None:
        seen: set[str] = set()
        dups: set[str] = set()
        for r in table.rows:
            a = extract_asin(r[asin_col] if asin_col < len(r) else None)
            if a is None:
                continue
            if a in seen:
                dups.add(a)
            seen.add(a)
        if dups:
            n = len(dups)
            warnings.append(
                f"{n} ASIN{'s appear' if n != 1 else ' appears'} more than once — only the first "
                f"row for each will be checked."
            )

    preview = [[display_value(c) for c in r] for r in table.rows[:PREVIEW_ROWS]]
    return DetectionResult(
        table=table,
        mapping=mapping,
        confidence=confidence,
        options=options,
        warnings=warnings,
        preview=preview,
    )


def raise_if_unmapped(detection: DetectionResult) -> None:
    """Raise a specific, friendly InputValidationError if the ASIN or expected
    price column couldn't be detected."""
    table = detection.table
    cols = ", ".join(f"'{h}'" for h in table.headers)
    where = f" (sheet '{table.sheet_name}')" if table.sheet_name else ""
    if detection.mapping.get("asin") is None:
        raise InputValidationError(
            f"We couldn't find a column with ASINs in your file{where}. Make sure one column has "
            f"ASINs (like B09W9FND7M) or Amazon product links. Columns found: {cols}."
        )
    if detection.mapping.get("expected_price") is None:
        raise InputValidationError(
            f"We couldn't find the expected price column in your file{where}. Make sure one column "
            f"has the price you expect (numbers like 1499), ideally headed 'Expected Price', "
            f"'Selling Price' or 'SP'. Columns found: {cols}."
        )


# ── Row parsing ─────────────────────────────────────────────────────────────

def _find_pincode_column(table: RawTable, taken: set[int]) -> Optional[int]:
    if not table.header_row:
        return None
    for j, h in enumerate(table.headers):
        if j not in taken and _normalize_header(h) in _PINCODE_HEADERS:
            return j
    return None


def _check_index(table: RawTable, idx, field_name: str) -> Optional[int]:
    if idx is None:
        return None
    if not isinstance(idx, int) or isinstance(idx, bool) or not 0 <= idx < len(table.headers):
        raise InputValidationError(f"The column chosen for {field_name} doesn't exist in this file.")
    return idx


def parse_rows(table: RawTable, mapping: dict[str, Optional[int]]) -> ParseReport:
    """Turn `table` into a ParseReport using `mapping` (field -> column index).

    `asin` and `expected_price` must be mapped; `brand` (and the unused
    `pincode`) are optional. Row-level problems — bad ASIN, missing or
    non-positive price, duplicate ASIN — mark the row invalid with a
    vendor-friendly reason; they never raise."""
    asin_col = _check_index(table, mapping.get("asin"), "the ASIN")
    price_col = _check_index(table, mapping.get("expected_price"), "the expected price")
    brand_col = _check_index(table, mapping.get("brand"), "the brand")
    if asin_col is None:
        raise InputValidationError("Please choose which column holds the ASINs.")
    if price_col is None:
        raise InputValidationError("Please choose which column holds the expected price.")
    if asin_col == price_col:
        raise InputValidationError("The ASIN and the expected price can't come from the same column.")
    if "pincode" in mapping:
        pincode_col = _check_index(table, mapping.get("pincode"), "the pincode")
    else:
        pincode_col = _find_pincode_column(table, {asin_col, price_col, brand_col} - {None})

    report = ParseReport(total_rows=len(table.rows), columns_found=list(table.headers))
    first_seen: dict[str, int] = {}

    for i, raw in enumerate(table.rows):
        row_number = table.row_number(i)

        def cell(idx: Optional[int]):
            return raw[idx] if idx is not None and idx < len(raw) else None

        raw_asin = cell(asin_col)
        raw_price = cell(price_col)
        asin = extract_asin(raw_asin)
        price = parse_price(raw_price)
        brand = display_value(cell(brand_col)) or None
        pincode = display_value(cell(pincode_col)) or None

        if asin is None:
            report.rows.append(ParsedRow(row_number, display_value(raw_asin), price, brand, pincode,
                                         False, _asin_error_reason(raw_asin)))
            continue
        if price is None:
            report.rows.append(ParsedRow(row_number, asin, None, brand, pincode,
                                         False, _price_error_reason(raw_price)))
            continue
        if asin in first_seen:
            report.rows.append(ParsedRow(row_number, asin, price, brand, pincode, False,
                                         f"Duplicate of row {first_seen[asin]} (same ASIN listed twice)"))
            continue
        first_seen[asin] = row_number
        report.rows.append(ParsedRow(row_number, asin, price, brand, pincode, True))

    report.valid_rows = len(report.valid)
    report.invalid_rows = len(report.invalid)
    return report
