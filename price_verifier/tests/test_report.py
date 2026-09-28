"""Offline Excel report tests (temp DB + temp output dir per test). Run:
  python -m unittest price_verifier.tests.test_report
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from price_verifier.excel.report import _sanitize_sheet_name, build_brand_report, build_report
from price_verifier.storage import checkpoint
from price_verifier.storage.db import init_db


class SheetNameSanitizationTests(unittest.TestCase):
    def test_invalid_characters_replaced(self):
        used: set[str] = set()
        name = _sanitize_sheet_name("A/V:Tech[India]", used)
        self.assertNotIn("/", name)
        self.assertNotIn(":", name)
        self.assertNotIn("[", name)
        self.assertNotIn("]", name)

    def test_truncated_to_31_chars(self):
        used: set[str] = set()
        name = _sanitize_sheet_name("A" * 50, used)
        self.assertLessEqual(len(name), 31)

    def test_collision_gets_unique_suffix(self):
        used: set[str] = set()
        first = _sanitize_sheet_name("Lapcare", used)
        second = _sanitize_sheet_name("Lapcare", used)
        self.assertNotEqual(first, second)

    def test_blank_brand_falls_back_to_placeholder(self):
        used: set[str] = set()
        name = _sanitize_sheet_name("", used)
        self.assertTrue(name)


class BuildReportTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmpdir.name) / "test.db"
        self.out_dir = Path(self.tmpdir.name) / "out"
        init_db(self.db_path)

    def tearDown(self):
        self.tmpdir.cleanup()

    def _seed_run(self) -> str:
        run_cfg = checkpoint.RunConfig(input_filename="t.csv")
        items = [
            checkpoint.RunItemRow(asin="B0000000A1", expected_price=1499.0, brand="Lapcare"),
            checkpoint.RunItemRow(asin="B0000000A2", expected_price=999.0, brand="Lapcare"),
            checkpoint.RunItemRow(asin="B0000000B1", expected_price=2999.0, brand="Portronics"),
            checkpoint.RunItemRow(asin="B0000000B2", expected_price=799.0, brand="Portronics"),
        ]
        run_id = checkpoint.create_run(run_cfg, items, db_path=self.db_path)
        checkpoint.mark_item_result(run_id, "B0000000A1", checkpoint.STATUS_MATCHED, actual_price=1499.0, db_path=self.db_path)
        checkpoint.mark_item_result(
            run_id, "B0000000A2", checkpoint.STATUS_MISMATCHED, actual_price=1050.0, mrp=1200.0,
            seller="Coco Blue Retail", product_title="Lapcare USB Hub", db_path=self.db_path,
        )
        checkpoint.mark_item_result(run_id, "B0000000B1", checkpoint.STATUS_MATCHED, actual_price=2999.0, db_path=self.db_path)
        checkpoint.mark_item_result(run_id, "B0000000B2", checkpoint.STATUS_FAILED, error_reason="captcha_challenge", db_path=self.db_path)
        checkpoint.finish_run(run_id, output_path=None, db_path=self.db_path)
        return run_id

    def test_one_sheet_per_brand_plus_overview_and_failed(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Overview", "Lapcare", "Portronics", "Could Not Verify"])

    def test_brand_sheet_contains_only_issue_rows(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = list(wb["Lapcare"].iter_rows(values_only=True))
        # header + exactly the one mismatched row (the matched B0000000A1 is excluded)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "B0000000A2")

    def test_clean_brand_sheet_is_header_only(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = list(wb["Portronics"].iter_rows(values_only=True))
        # B0000000B1 matched (excluded), B0000000B2 failed (goes to Could Not Verify, not here)
        self.assertEqual(len(rows), 1)

    def test_failed_rows_land_in_could_not_verify_not_brand_sheet(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = list(wb["Could Not Verify"].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "B0000000B2")

    def test_difference_column_computed(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = list(wb["Lapcare"].iter_rows(values_only=True))
        # 1050 actual - 999 expected = 51
        self.assertEqual(rows[1][6], 51)

    def test_seller_and_mrp_present_in_issue_row(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = list(wb["Lapcare"].iter_rows(values_only=True))
        self.assertEqual(rows[1][2], "Coco Blue Retail")  # Seller column
        self.assertEqual(rows[1][5], 1200.0)  # MRP column

    def test_overview_counts_per_brand(self):
        run_id = self._seed_run()
        path = build_report(run_id, output_path=self.out_dir / "r.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = {r[0]: r for r in wb["Overview"].iter_rows(values_only=True, min_row=2)}
        self.assertEqual(rows["Lapcare"][1], 2)  # total
        self.assertEqual(rows["Lapcare"][2], 1)  # matched
        self.assertEqual(rows["Lapcare"][3], 1)  # mismatched

    def test_build_brand_report_single_brand_only(self):
        run_id = self._seed_run()
        path = build_brand_report(run_id, "Lapcare", output_path=self.out_dir / "brand.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Lapcare"])
        rows = list(wb["Lapcare"].iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)  # header + 1 issue row

    def test_build_brand_report_for_clean_brand_is_empty(self):
        run_id = self._seed_run()
        path = build_brand_report(run_id, "Portronics", output_path=self.out_dir / "brand2.xlsx", db_path=self.db_path)
        wb = load_workbook(path)
        rows = list(wb["Portronics"].iter_rows(values_only=True))
        self.assertEqual(len(rows), 1)  # header only — no issues, no failures shown here


if __name__ == "__main__":
    unittest.main()
