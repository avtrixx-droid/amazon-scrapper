"""
e2e_frozen.py — drive the BUILT app (the Windows .exe in CI, or the Linux
binary from the same spec) through the whole user journey, as a real
separate process, exactly as the vendor uses it:

    activation (wrong key refused, right key accepted)
    -> upload -> automatic column mapping -> confirm -> run
       (including rows only a real Chrome can price: the Chrome check runs)
    -> results -> Excel download (checked) -> history

and on every step: the app process is still alive and every request is
answered by the SAME process — the failure mode behind "Lost connection to
the progress feed", where a second copy of the app shared the port.

Not a unit test (no test_ prefix). Needs a build whose license server URL
points at this script's fake license server:

    python -c "open('price_verifier/_build_config.py','w').write('SERVER_URL = \\'http://127.0.0.1:8123\\'\\n')"
    pyinstaller price_verifier_windows.spec --distpath dist-e2e
    python -m price_verifier.tests.e2e_frozen --exe dist-e2e/PriceVerificationTool.exe --expect-chrome

Exit code 0 = every check passed. The app's logs/ (and data/debug_html/)
are copied to --artifacts for inspection either way.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import httpx
from openpyxl import load_workbook

from price_verifier.tests.sim_amazon import SimAmazon, ThrottlePolicy, make_catalog

E2E_KEY = "AMZ-E2E2-E2E2-E2E2-E2E2"
APP = "http://127.0.0.1:5001"


class FakeLicenseServer:
    """Accepts exactly E2E_KEY for product price_verifier; records calls."""

    def __init__(self, port: int):
        self.calls: list[tuple[str, dict]] = []
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_GET(self):
                self._reply(200, {"ok": True})

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers.get("Content-Length") or 0)) or b"{}")
                outer.calls.append((self.path, body))
                if body.get("key") != E2E_KEY:
                    return self._reply(404, {"ok": False, "reason": "key_not_found"})
                if body.get("product") != "price_verifier":
                    return self._reply(403, {"ok": False, "reason": "product_not_licensed"})
                self._reply(200, {"ok": True, "customer": "E2E Test", "expires_at": "2099-01-01T00:00:00Z",
                                  "run_token": "t", "products": ["price_verifier"]})

            def _reply(self, code, data):
                raw = json.dumps(data).encode()
                self.send_response(code)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)

        self.server = ThreadingHTTPServer(("127.0.0.1", port), Handler)
        self.server.daemon_threads = True
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def stop(self):
        self.server.shutdown()


class Failed(Exception):
    pass


class Journey:
    def __init__(self, args):
        self.args = args
        self.exe = Path(args.exe).resolve()
        self.proc: subprocess.Popen | None = None
        self.pids: set[str] = set()
        self.log: list[str] = []
        self.run_id: str | None = None
        self.engine_killed = False
        self.http = httpx.Client(base_url=APP, follow_redirects=False, trust_env=False, timeout=90)

    # ── helpers ────────────────────────────────────────────────────────────
    def say(self, msg):
        line = f"[{time.strftime('%H:%M:%S')}] {msg}"
        self.log.append(line)
        print(line, flush=True)

    def check(self, cond, msg):
        if not cond:
            raise Failed(msg)
        self.say(f"OK  {msg}")

    def alive(self):
        if self.proc is not None and self.proc.poll() is not None:
            raise Failed(f"the app process EXITED (code {self.proc.returncode}) — see logs/crash.log, logs/app.log")
        r = self.http.get("/healthz")
        m = re.search(r"pid=(\d+)", r.text)
        if not m:
            raise Failed(f"unexpected /healthz answer: {r.text[:80]!r}")
        self.pids.add(m.group(1))
        if len(self.pids) > 1:
            raise Failed(f"requests answered by MORE THAN ONE app process: pids {sorted(self.pids)}")

    # ── journey ────────────────────────────────────────────────────────────
    def launch(self, sim_url):
        env = dict(os.environ, PV_NO_BROWSER="1", NO_PROXY="127.0.0.1,localhost",
                   no_proxy="127.0.0.1,localhost")
        if sim_url:
            env["PV_MARKETPLACE_BASE_URL"] = sim_url
        else:
            env.pop("PV_MARKETPLACE_BASE_URL", None)
        kwargs = {}
        if os.name == "nt":
            kwargs["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
        else:
            kwargs["start_new_session"] = True
        self.proc = subprocess.Popen([str(self.exe)], cwd=str(self.exe.parent), env=env, **kwargs)
        deadline = time.time() + 90
        while time.time() < deadline:
            try:
                self.alive()
                self.say(f"app is up (server pid {next(iter(self.pids))})")
                return
            except (httpx.HTTPError, Failed) as e:
                if isinstance(e, Failed) and "EXITED" in str(e):
                    raise
                time.sleep(1)
        raise Failed("the app did not start within 90 s")

    def activation(self):
        r = self.http.get("/")
        self.check(r.status_code == 302 and r.headers["location"].endswith("/activate"),
                   "without a license every page leads to Activate")
        r = self.http.post("/activate", data={"key": "AMZ-WRNG-WRNG-WRNG-WRNG"})
        self.check(r.status_code == 200 and "find that license key" in r.text, "a wrong key is refused")
        r = self.http.post("/activate", data={"key": E2E_KEY.lower()})
        self.check(r.status_code == 302, "the right key activates")
        r = self.http.get("/")
        self.check(r.status_code == 200 and "Licensed to E2E Test" in r.text, "upload page opens, shows licensee")
        self.alive()

    def run(self, csv_path, n_rows):
        r = self.http.post("/upload", files={"file": ("vendor list.csv", open(csv_path, "rb"), "text/csv")})
        self.check(r.status_code == 302 and "/map/" in r.headers["location"], "upload accepted")
        map_url = r.headers["location"]
        page = self.http.get(map_url).text
        sel = dict(re.findall(r'name="col_(\w+)".*?<option value="(\d*)"[^>]*selected', page, re.S))
        self.check(sel.get("asin") == "0" and sel.get("expected_price") == "2" and sel.get("brand") == "1",
                   f"columns detected automatically {sel}")
        r = self.http.post(map_url, data={f"col_{k}": v for k, v in sel.items()})
        self.check(f"{n_rows}</div>" in r.text.replace(" ", "") or "Ready to check" in r.text, "confirm page")
        upload_id = map_url.rstrip("/").rsplit("/", 1)[-1]
        r = self.http.post("/start", data={"upload_id": upload_id, "concurrency": "15", "tolerance_abs": "1",
                                           "tolerance_pct": "0", "use_browser": "1"})
        self.check(r.status_code == 302 and "/progress/" in r.headers["location"], "run started")
        run_id = r.headers["location"].rsplit("/", 1)[-1]
        self.run_id = run_id
        self.check(self.http.get(f"/progress/{run_id}").status_code == 200, "progress page opens")

        t0, last, deadline = time.time(), None, time.time() + self.args.timeout
        while True:
            # Re-open the live feed whenever it ends early, like the page does,
            # and probe the app between reads.
            try:
                with self.http.stream("GET", f"/stream/{run_id}", timeout=httpx.Timeout(30, read=30)) as s:
                    for line in s.iter_lines():
                        if line.startswith("data:"):
                            last = json.loads(line[5:])
                            if last.get("status") != "running":
                                break
                            if int(time.time() - t0) % 15 == 0:
                                self.say(f"    {last['done']}/{last['total']}  {last.get('phase_label')}")
                            if (self.args.kill_engine_at and not self.engine_killed
                                    and last["done"] >= self.args.kill_engine_at):
                                self.kill_engine()
            except httpx.ReadTimeout:
                pass
            self.alive()
            if last and last.get("status") != "running":
                break
            if time.time() > deadline:
                raise Failed(f"run did not finish within {self.args.timeout} s (last: {last})")
        self.say(f"run finished in {time.time() - t0:.0f} s: {json.dumps(last)[:200]}")
        self.check(last["status"] == "completed", "run completed")
        if self.args.kill_engine_at:
            self.check(self.engine_killed, "the engine was killed mid-run and the run still completed")
        self.check(last["done"] == n_rows, f"all {n_rows} rows have a result")
        if self.args.expect_chrome and not self.args.live:
            self.check(last["failed"] == 0, "no row left as Could Not Verify")
        return run_id, last

    def outputs(self, run_id):
        r = self.http.get(f"/results/{run_id}")
        self.check(r.status_code == 200 and "Download" in r.text, "results page")
        r = self.http.get(f"/download/{run_id}")
        self.check(r.status_code == 200 and r.content[:2] == b"PK", "Excel report downloads")
        wb = load_workbook(io.BytesIO(r.content))
        self.check({"Overview", "Could Not Verify", "Run Info"} <= set(wb.sheetnames), f"report sheets {wb.sheetnames}")
        info = {str(row[0].value): row[1].value for row in wb["Run Info"].iter_rows(min_row=2) if row[0].value}
        chrome_rows = int(info.get("Resolved via Google Chrome check") or 0)
        self.say(f"    Run Info: first pass={info.get('Resolved on first pass')}, Chrome={chrome_rows}, "
                 f"could not verify={info.get('Could not verify')}")
        if self.args.expect_chrome and not self.args.live:
            self.check(chrome_rows >= self.args.browser_only,
                       f"the Chrome check priced the {self.args.browser_only} browser-only rows")
        self.check(self.http.get("/history").status_code == 200, "history page")
        for _ in range(20):
            self.alive()
        self.check(len(self.pids) == 1, "every request answered by one and the same app process")

    def kill_engine(self):
        """Kill the scraping engine process(es) hard, the way a native crash
        ends them — the app must stay up and finish the run on a new one."""
        import psutil

        app = psutil.Process(self.proc.pid)
        engines = [p for p in app.children(recursive=True)
                   if "--multiprocessing-fork" in " ".join(p.cmdline())]
        self.check(bool(engines), f"found the engine process to kill (pids {[p.pid for p in engines]})")
        for p in engines:
            p.kill()
        self.engine_killed = True
        self.say(f"    killed engine process(es) {[p.pid for p in engines]} mid-run")

    def diagnose(self):
        """Print why it failed into the job output itself (the uploaded logs
        artifact isn't always reachable): the Could Not Verify reasons, then
        every warning/error and traceback the app logged."""
        print("=" * 30 + " DIAGNOSTICS " + "=" * 30, flush=True)
        if self.run_id:
            try:
                r = self.http.get(f"/download/{self.run_id}")
                wb = load_workbook(io.BytesIO(r.content), read_only=True)
                if "Run Info" in wb.sheetnames:
                    print("-- Run Info sheet --")
                    for row in wb["Run Info"].iter_rows(values_only=True):
                        if any(v is not None for v in row):
                            print("   ", " | ".join("" if v is None else str(v) for v in row))
                if "Could Not Verify" in wb.sheetnames:
                    print("-- Could Not Verify sheet --")
                    for row in wb["Could Not Verify"].iter_rows(values_only=True):
                        print("   ", " | ".join("" if v is None else str(v) for v in row))
            except Exception as e:  # noqa: BLE001
                print(f"(report not available: {type(e).__name__}: {e})")
        logs = self.exe.parent / "logs"
        for name in ("app.log", "crash.log", "console.log", "startup.log"):
            f = logs / name
            if not f.exists() or not f.stat().st_size:
                continue
            lines = f.read_text(encoding="utf-8", errors="replace").splitlines()
            if name == "app.log" and not self.args.live:
                keep, in_tb = [], False
                for ln in lines:
                    if re.search(r"\b(WARNING|ERROR|CRITICAL)\b", ln):
                        keep.append(ln)
                        in_tb = False
                    elif ln.startswith(("Traceback", "  ", "\t")) or in_tb:
                        keep.append(ln)
                        in_tb = not re.match(r"^\d{4}-\d\d-\d\d", ln)
                lines = keep
            print(f"-- logs/{name} ({len(lines)} lines shown, last 400) --")
            for ln in lines[-400:]:
                print("   ", ln)
        if self.args.live:
            self.dump_debug_pages()
        print("=" * 73, flush=True)

    def dump_debug_pages(self):
        """A live run's unrecognised pages (and offers-page samples), printed
        gzip+base64 between markers so they can become parser fixtures."""
        import base64
        import gzip

        root = self.exe.parent / "data" / "debug_html"
        if not root.exists():
            return
        files = sorted(root.rglob("*.html"))
        picks = ([f for f in files if "unrecognised" in f.name][:2]
                 + [f for f in files if "offers-sample" in f.name][:2])
        for f in picks:
            blob = base64.b64encode(gzip.compress(f.read_bytes(), 9)).decode()
            print(f"=====BEGIN PAGE {f.stem[:80]}=====")
            for i in range(0, len(blob), 4000):
                print(blob[i:i + 4000])
            print(f"=====END PAGE {f.stem[:80]}=====", flush=True)

    def stop(self):
        if self.proc is None or self.proc.poll() is not None:
            return
        if os.name == "nt":
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(self.proc.pid)], capture_output=True)
        else:
            try:
                os.killpg(self.proc.pid, signal.SIGTERM)
            except OSError:
                pass
        try:
            self.proc.wait(timeout=20)
        except subprocess.TimeoutExpired:
            self.proc.kill()


_FALLBACK_LIVE_ASINS = (
    # Used only if the search page can't be read (e.g. the runner is blocked).
    "B0BSHF7WHW", "B0CHX1W1XY", "B0D1XD1ZV3", "B09G9FPHY6", "B07WFPMPX3", "B08L5WD9D6",
)


def harvest_live_asins(n: int, query: str) -> list[str]:
    """Real ASINs from amazon.in searches (through Chrome — plain HTTP from a
    CI address mostly gets Amazon's bot-check page), so the live journey
    checks real product pages. Falls back to a fixed list."""
    found: list[str] = []
    try:
        from price_verifier.tests.live_probe import chrome_harvest

        found, _ = chrome_harvest([q.strip() for q in query.split(",") if q.strip()], n,
                                  Path(tempfile.mkdtemp(prefix="pv_harvest_")), fetch=False)
    except Exception as e:  # noqa: BLE001
        print(f"(could not harvest ASINs: {type(e).__name__}: {e})", flush=True)
    for a in _FALLBACK_LIVE_ASINS:
        if len(found) >= n:
            break
        if a not in found:
            found.append(a)
    return found[:n]


def main() -> int:
    for stream in (sys.stdout, sys.stderr):   # the Windows CI console is cp1252; reports contain ₹
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="End-to-end check of the built Price Verification Tool")
    ap.add_argument("--exe", required=True)
    ap.add_argument("--license-port", type=int, default=8123)
    ap.add_argument("--asins", type=int, default=24)
    ap.add_argument("--browser-only", type=int, default=3)
    ap.add_argument("--expect-chrome", action="store_true",
                    help="Chrome is installed: require the Chrome check to price the browser-only rows")
    ap.add_argument("--timeout", type=int, default=1200)
    ap.add_argument("--artifacts", default="e2e-artifacts")
    ap.add_argument("--kill-engine-at", type=int, default=0,
                    help="kill the engine process once this many rows are done (needs psutil)")
    ap.add_argument("--live", action="store_true",
                    help="real amazon.in instead of the simulator (needs internet): checks the app "
                         "survives a real run end to end; doesn't require every row to be priced")
    ap.add_argument("--live-query", default="lapcare,lapcare keyboard,lapcare mouse,lapcare webcam,lapcare charger,lapcare laptop adapter,lapcare headphones,lapcare cable,lapcare speaker")
    args = ap.parse_args()

    lic = FakeLicenseServer(args.license_port)
    tmp = Path(tempfile.mkdtemp(prefix="pv_e2e_"))
    csv_path = tmp / "vendor list.csv"
    sim = None
    if args.live:
        live_asins = harvest_live_asins(args.asins, args.live_query)
        n_rows = len(live_asins)
        with open(csv_path, "w", encoding="utf-8") as f:
            f.write("ASIN No.,Brand Name,SP\n")
            for i, a in enumerate(live_asins):
                f.write(f"{a},Brand {i % 3},{499 + i}\n")
    else:
        catalog = make_catalog(args.asins, seed=21, kinds={"unavailable": 0.05, "no_offer": 0.05})
        n_rows = len(catalog)
        browser_only = [a for a, p in catalog.items() if p.kind == "in_stock"][:args.browser_only]
        # Offers page in a layout the tool won't trust, so those rows really go to Chrome.
        sim = SimAmazon(catalog, ThrottlePolicy(burst=60, sustained_rps=10, anon_burst=30, ip_ceiling_rps=30),
                        padding_kb=80, aod_layout="changed", browser_only=set(browser_only)).start()
        with open(csv_path, "w", encoding="utf-8") as f:
            f.write("ASIN No.,Brand Name,SP\n")
            for i, p in enumerate(catalog.values()):
                f.write(f"{p.asin},{p.brand},{(p.price if i % 5 else p.price + 20):.0f}\n")

    # Every journey starts unlicensed (it checks activation): remove a
    # license file an earlier journey on this machine left behind.
    try:
        from price_verifier import licensing

        licensing.client().license_path.unlink(missing_ok=True)
    except Exception as e:  # noqa: BLE001
        print(f"(could not clear the license file: {e})", flush=True)

    j = Journey(args)
    ok = False
    try:
        j.launch(sim.base_url if sim else None)
        j.activation()
        run_id, _ = j.run(csv_path, n_rows)
        j.outputs(run_id)
        runs = [b for path, b in lic.calls if path == "/authorize-run"]
        j.check(len(runs) == 1 and runs[0].get("asin_count") == n_rows and runs[0].get("product") == "price_verifier",
                "the license server authorized exactly this run, for price_verifier")
        ok = True
        j.say("ALL END-TO-END CHECKS PASSED")
    except Failed as e:
        j.say(f"FAIL {e}")
    except Exception as e:  # noqa: BLE001 — report anything as a failure with context
        j.say(f"FAIL unexpected {type(e).__name__}: {e}")
    finally:
        if not ok or args.live:   # a live run always shows what really happened
            try:
                j.diagnose()
            except Exception as e:  # noqa: BLE001
                print(f"(diagnostics failed: {type(e).__name__}: {e})")
        j.stop()
        if sim is not None:
            sim.stop()
        lic.stop()
        out = Path(args.artifacts)
        out.mkdir(parents=True, exist_ok=True)
        (out / "e2e.log").write_text("\n".join(j.log), encoding="utf-8")
        for sub in ("logs", "data/debug_html"):
            src = j.exe.parent / sub
            if src.exists():
                shutil.copytree(src, out / sub.replace("/", "_"), dirs_exist_ok=True)
        print(f"artifacts in {out.resolve()}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
