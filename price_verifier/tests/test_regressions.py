"""
test_regressions.py — one test per defect found in the independent review of
v2, so each stays fixed:

  * expected-price column: MRP / GST-exclusive / old-price columns must not
    be pre-selected over the real selling price
  * "₹" mangled to "?" by Excel's ANSI CSV export still parses
  * ₹1 boundary is exact (no floating-point false flags)
  * "Lapcare" / "LAPCARE" are one brand (one sheet, one download)
  * a Chrome start that never returns can't hang the run
  * a report file left open in Excel doesn't break the rebuild
"""

from __future__ import annotations

import asyncio
import io
import tempfile
import threading
import time
import os
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from price_verifier.excel import report
from price_verifier.ingest.column_detect import detect_columns, load_table, parse_price
from price_verifier.pipeline.compare import prices_match
from price_verifier.storage import checkpoint
from price_verifier.storage.db import init_db
from price_verifier.tests.test_pipeline_integration import FakeAmazon, PipelineHarness, asins, fast_tuning

ASINS = [f"B0{100000 + i}ZZ" for i in range(20)]


def _pick_price_column(headers: list[str], row: list) -> str:
    wb = Workbook()
    ws = wb.active
    ws.append(["ASIN"] + headers)
    for a in ASINS:
        ws.append([a] + row)
    buf = io.BytesIO()
    wb.save(buf)
    table = load_table("list.xlsx", buf.getvalue())
    return table.headers[detect_columns(table).mapping["expected_price"]]


class ExpectedPriceColumnTests(unittest.TestCase):
    CASES = [
        (["MRP", "Price after Discount"], [1999, 1499], "Price after Discount"),
        (["MRP", "Discount Price"], [1999, 1499], "Discount Price"),
        (["MRP", "Price (incl. GST)"], [1999, 1499], "Price (incl. GST)"),
        (["MRP", "Price w/o GST", "Price with GST"], [1999, 1270, 1499], "Price with GST"),
        (["MRP", "Price (after 10% off)"], [1999, 1499], "Price (after 10% off)"),
        (["MRP", "SP excl GST", "SP incl GST"], [1999, 1270, 1499], "SP incl GST"),
        (["Old Price", "New Price"], [1599, 1499], "New Price"),
        (["Current SP", "Revised SP"], [1599, 1499], "Revised SP"),
        (["MRP", "Selling Price", "Discount %"], [1999, 1499, 25], "Selling Price"),
        (["Offer %", "Price"], [10, 1499], "Price"),
        (["Qty", "MRP", "Offer Price", "GST %"], [5, 1999, 1499, 18], "Offer Price"),
        (["Cost Price", "MRP", "Selling Price (Rs.)"], [900, 1999, 1499], "Selling Price (Rs.)"),
    ]

    def test_selling_price_column_is_preselected(self):
        for headers, row, want in self.CASES:
            with self.subTest(headers=headers):
                self.assertEqual(_pick_price_column(headers, row), want)

    def test_mrp_alone_is_still_used_as_last_resort(self):
        self.assertEqual(_pick_price_column(["MRP"], [1999]), "MRP")


class MangledRupeeTests(unittest.TestCase):
    def test_rupee_sign_mangled_by_ansi_export(self):
        for v in ("?1,499", "? 1,499.00", "�1499", "â‚¹1,499"):
            with self.subTest(v=v):
                self.assertEqual(parse_price(v), 1499.0)
        self.assertIsNone(parse_price("?"))

    def test_cp1252_csv_with_rupee_prices_is_readable(self):
        text = "ASIN,Brand,SP\r\n" + "".join(f'{a},Lapcare,"₹1,{i:03d}"\r\n' for i, a in enumerate(ASINS))
        table = load_table("list.csv", text.encode("cp1252", errors="replace"))
        det = detect_columns(table)
        self.assertEqual(table.headers[det.mapping["expected_price"]], "SP")


class ToleranceBoundaryTests(unittest.TestCase):
    def test_exactly_one_rupee_is_never_flagged(self):
        misflagged = [p for p in range(100, 500_000, 7) if not prices_match(p / 100, (p + 100) / 100, 1.0, 0.0)]
        self.assertEqual(misflagged, [])
        self.assertFalse(prices_match(1499.0, 1500.01, 1.0, 0.0))


class BrandCaseTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmpdir.name) / "t.db"
        init_db(self.db_path)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_brand_spellings_are_merged(self):
        items = [
            checkpoint.RunItemRow(asin="B0000000A1", expected_price=100.0, brand="Lapcare"),
            checkpoint.RunItemRow(asin="B0000000A2", expected_price=100.0, brand="Lapcare"),
            checkpoint.RunItemRow(asin="B0000000A3", expected_price=100.0, brand="LAPCARE"),
            checkpoint.RunItemRow(asin="B0000000A4", expected_price=100.0, brand=""),
        ]
        run_id = checkpoint.create_run(checkpoint.RunConfig(input_filename="t.csv"), items, db_path=self.db_path)
        for a in ("B0000000A1", "B0000000A3"):
            checkpoint.mark_item_result(run_id, a, checkpoint.STATUS_MISMATCHED, actual_price=150.0,
                                        db_path=self.db_path)
        checkpoint.mark_item_result(run_id, "B0000000A4", checkpoint.STATUS_MISMATCHED, actual_price=150.0,
                                    scraped_brand="lapcare ", db_path=self.db_path)
        self.assertEqual(checkpoint.get_brand_issue_counts(run_id, db_path=self.db_path),
                         [{"brand": "Lapcare", "issues": 3}])
        out = Path(self.tmpdir.name) / "r.xlsx"
        wb = load_workbook(report.build_report(run_id, output_path=out, db_path=self.db_path))
        self.assertEqual([n for n in wb.sheetnames if "apcare" in n.lower()], ["Lapcare"])


class ReportLockedTests(unittest.TestCase):
    def test_locked_output_falls_back_to_new_name(self):
        with tempfile.TemporaryDirectory() as d:
            target = Path(d) / "PriceVerification_x.xlsx"
            real_save = Workbook.save
            calls = []

            def save(wb, filename):
                calls.append(filename)
                if filename == str(target):
                    raise PermissionError("file is open in Excel")
                return real_save(wb, filename)

            with mock.patch.object(Workbook, "save", save):
                path = report._save_workbook(Workbook(), target)
            self.assertNotEqual(path, target)
            self.assertTrue(path.exists())
            self.assertTrue(path.name.startswith("PriceVerification_x_"))


class StuckChromeStartTests(PipelineHarness):
    async def test_chrome_start_that_never_returns_cannot_hang_the_run(self):
        release = threading.Event()

        class StuckBrowser:
            closed = 0

            def start(self):
                release.wait(30)  # e.g. ChromeDriver download through a stalling proxy

            def close(self):
                StuckBrowser.closed += 1

        amazon = FakeAmazon(ambiguous=set(asins(3)))
        run_id, pending = self.make_run(asins(3))
        t0 = time.monotonic()
        stats = await self.go(run_id, pending, amazon, StuckBrowser(),
                              tuning=fast_tuning(browser_start_timeout=0.3))
        self.assertLess(time.monotonic() - t0, 5.0)
        self.assertEqual(stats.failed, 3)
        rows = self.rows(run_id)
        self.assertTrue(all("took too long to start" in r["error_reason"] for r in rows.values()))
        # when the stuck start finally returns, the late Chrome is closed, not orphaned
        release.set()
        for _ in range(50):
            if StuckBrowser.closed:
                break
            await asyncio.sleep(0.02)
        self.assertGreaterEqual(StuckBrowser.closed, 1)


if __name__ == "__main__":
    unittest.main()


class ConcurrentCountTests(unittest.TestCase):
    """Rows finishing at the same moment (many workers) must never leave
    the run's totals stale — the results page's Retry button and the
    history page read them."""

    def test_parallel_row_writes_keep_run_totals_exact(self):
        from concurrent.futures import ThreadPoolExecutor

        with tempfile.TemporaryDirectory() as d:
            db = Path(d) / "t.db"
            init_db(db)
            items = [checkpoint.RunItemRow(asin=f"B0CONC{i:04d}", expected_price=10.0) for i in range(60)]
            run_id = checkpoint.create_run(checkpoint.RunConfig(input_filename="t.csv"), items, db_path=db)
            # first every row fails, then a retry settles all of them — in parallel both times
            for status in (checkpoint.STATUS_FAILED, checkpoint.STATUS_MATCHED):
                with ThreadPoolExecutor(16) as pool:
                    list(pool.map(lambda it: checkpoint.mark_item_result(
                        run_id, it.asin, status, actual_price=10.0, db_path=db), items))
            run = checkpoint.get_run(run_id, db_path=db)
            self.assertEqual((run["matched"], run["failed"]), (60, 0))


class FrozenWindowsRuntimeTests(unittest.TestCase):
    """Safeguards for the windowed Windows .exe ("Lost connection to the
    progress feed"): a second copy of the app — a multiprocessing child
    without freeze_support(), or the .exe started again while the first,
    windowless copy still runs — must never become a second server on the
    same port (Windows allows it because Werkzeug sets SO_REUSEADDR)."""

    def test_freeze_support_runs_before_anything_else(self):
        src = (Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8")
        block = src[src.index('if __name__ == "__main__":'):]
        self.assertLess(block.index("multiprocessing.freeze_support()"), block.index("main()"))

    def test_port_is_bound_exclusively_on_windows(self):
        import werkzeug.serving as ws

        from price_verifier import app as appmod

        saved = (ws.BaseWSGIServer.allow_reuse_address, ws.BaseWSGIServer.server_bind,
                 getattr(ws.BaseWSGIServer, "_pv_exclusive", None))
        try:
            with mock.patch.object(appmod.sys, "platform", "win32"):
                appmod._exclusive_port_on_windows()
            self.assertFalse(ws.BaseWSGIServer.allow_reuse_address)
            self.assertTrue(ws.BaseWSGIServer._pv_exclusive)
        finally:
            ws.BaseWSGIServer.allow_reuse_address, ws.BaseWSGIServer.server_bind = saved[0], saved[1]
            if saved[2] is None:
                del ws.BaseWSGIServer._pv_exclusive

    def test_stream_for_a_run_this_process_does_not_know(self):
        from price_verifier import app as appmod

        with mock.patch.dict(os.environ, {"PV_LICENSE_DISABLED": "1"}):
            r = appmod.app.test_client().get("/stream/no-such-run")
            self.assertIn('"status": "unknown"', r.get_data(as_text=True))

    def test_healthz_names_the_process(self):
        from price_verifier import app as appmod

        r = appmod.app.test_client().get("/healthz").get_data(as_text=True)
        self.assertIn("price-verifier-ok", r)
        self.assertIn(f"pid={os.getpid()}", r)


class ReadOnlyInstallFolderTests(unittest.TestCase):
    """The frozen .exe keeps its data next to itself — unless that folder is
    read-only (Program Files, a locked share): then the user's app-data
    folder, instead of dying at import before any error can be shown."""

    def test_frozen_base_dir_falls_back_when_exe_folder_is_read_only(self):
        from price_verifier import config

        with tempfile.TemporaryDirectory() as d:
            exe_dir, appdata = Path(d) / "ro", Path(d) / "appdata"
            exe_dir.mkdir()
            with mock.patch.object(config.sys, "frozen", True, create=True), \
                    mock.patch.object(config.sys, "executable", str(exe_dir / "PriceVerificationTool.exe")), \
                    mock.patch.dict(os.environ, {"LOCALAPPDATA": str(appdata)}):
                self.assertEqual(config._get_base_dir(), exe_dir)
                with mock.patch.object(config, "_writable", lambda folder: False):
                    self.assertEqual(config._get_base_dir(), appdata / "PriceVerificationTool")
                self.assertTrue((appdata / "PriceVerificationTool").is_dir())
