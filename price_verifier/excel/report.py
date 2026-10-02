"""
report.py — builds the output workbook.

Shaped around the vendor's workflow: tally live Amazon price against what
was agreed with a seller, and email that seller when it's wrong. So:

- One sheet PER BRAND (each brand's issues go to a different seller).
- Each brand sheet lists ONLY rows with an issue — price mismatch beyond
  tolerance, out of stock, unavailable, no buy box, ASIN not found.
- Columns follow the vendor's own description: ASIN, then expected price,
  then the price Amazon shows, the difference, MRP, seller — then context.
- Rows that couldn't be checked at all are not a pricing issue to send
  anyone; they go on one "Could Not Verify" sheet (and the app offers a
  Retry button for them). An "Overview" sheet summarises per brand and a
  "Run Info" sheet records how the run went.

Reads straight from SQLite, so History re-downloads and the per-brand
single-sheet export are just this module re-run against a past run_id.
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from price_verifier import config
from price_verifier.storage import checkpoint

HEADER_FILL = PatternFill("solid", fgColor="003366")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="F7F7F7")
MISMATCH_FILL = PatternFill("solid", fgColor="FFE0E0")
LISTING_FILL = PatternFill("solid", fgColor="FFF4D6")
LINK_FONT = Font(color="0563C1", underline="single")
MONEY_FORMAT = "#,##0.00"

_STATUS_NO_OFFER = getattr(checkpoint, "STATUS_NO_FEATURED_OFFER", "no_featured_offer")

_STATUS_DISPLAY = {
    checkpoint.STATUS_MISMATCHED: "PRICE MISMATCH",
    checkpoint.STATUS_OUT_OF_STOCK: "OUT OF STOCK",
    checkpoint.STATUS_UNAVAILABLE: "UNAVAILABLE",
    checkpoint.STATUS_NOT_FOUND: "ASIN NOT FOUND",
    _STATUS_NO_OFFER: "NO BUY BOX",
}

# Internal failure codes -> something a vendor can act on.
_FRIENDLY_REASONS = (
    ("soft_block", "Amazon temporarily limited requests — use Retry"),
    ("throttled", "Amazon temporarily limited requests — use Retry"),
    ("captcha", "Amazon asked for a CAPTCHA — use Retry"),
    ("timeout", "Amazon took too long to respond — use Retry"),
    ("connection", "Network problem reaching Amazon — check internet, then Retry"),
    ("chrome", "Final check in Google Chrome wasn't possible — install/update Chrome, then Retry"),
    ("browser", "Final check in Google Chrome failed — use Retry"),
    ("selector", "Page layout not recognised — saved for diagnosis; use Retry"),
    ("price could not be parsed", "Price not shown on the page — saved for diagnosis; use Retry"),
)

_INVALID_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")

# Excel/CSV formula-injection hardening: title/seller/brand come from scraped
# Amazon HTML (seller-controlled) or the vendor's upload, and this workbook
# gets forwarded by email. A leading =/+/-/@ is what spreadsheet apps key off
# to reinterpret text as a formula; an apostrophe prefix neutralises it.
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def _sanitize_cell_text(value):
    if isinstance(value, str) and value and value[0] in _FORMULA_TRIGGER_CHARS:
        return "'" + value
    return value


def _sanitize_sheet_name(name: str, used: set[str]) -> str:
    """Excel sheet names: <=31 chars, none of \\ / ? * [ ] :, unique
    (case-insensitively) within the workbook."""
    cleaned = _INVALID_SHEET_CHARS.sub(" ", (name or "Unknown Brand").strip()) or "Unknown Brand"
    cleaned = cleaned[:31]
    candidate = cleaned
    n = 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _local_time(iso: str | None) -> str | None:
    if not iso:
        return None
    try:
        return datetime.fromisoformat(iso).astimezone().strftime("%d %b %Y, %I:%M %p")
    except ValueError:
        return iso


def _friendly_reason(reason: str | None) -> str | None:
    if not reason:
        return None
    low = reason.lower()
    for needle, text in _FRIENDLY_REASONS:
        if needle in low:
            return text
    return reason


def _brand(r: dict) -> str:
    # get_run_items() has already merged case variants ("LAPCARE" / "Lapcare").
    return r.get("effective_brand") or checkpoint.effective_brand(r)


def _write_table(ws: Worksheet, headers: list[str], rows: list[tuple], widths: dict[int, int] | None = None) -> None:
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for i, row in enumerate(rows, start=2):
        ws.append(row)
        if i % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = ALT_FILL

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        if widths and col in widths:
            ws.column_dimensions[letter].width = widths[col]
            continue
        max_len = len(str(headers[col - 1]))
        for row in rows[:500]:
            val = row[col - 1] if col - 1 < len(row) else ""
            max_len = max(max_len, len(str(val)) if val is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)


_ISSUE_HEADERS = [
    "ASIN", "Expected Price (₹)", "Price on Amazon (₹)", "Difference (₹)", "MRP (₹)",
    "Seller", "Product Title", "Status", "Note", "Amazon Link", "Checked At",
]
_ISSUE_MONEY_COLS = (2, 3, 4, 5)
_ISSUE_LINK_COL = 10


def _issue_note(r: dict) -> str | None:
    if r["status"] == checkpoint.STATUS_MISMATCHED and r["actual_price"] is not None:
        diff = r["actual_price"] - r["expected_price"]
        direction = "higher" if diff > 0 else "lower"
        return f"Amazon shows ₹{abs(diff):,.2f} {direction} than expected"
    if r["status"] == _STATUS_NO_OFFER:
        return "No seller holds the buy box (Amazon shows 'See All Buying Options')"
    return _friendly_reason(r.get("error_reason"))


def _issue_row(r: dict) -> tuple:
    expected = r["expected_price"]
    actual = r["actual_price"]
    diff = round(actual - expected, 2) if actual is not None else None
    return (
        r["asin"],
        expected,
        actual,
        diff,
        r.get("mrp"),
        _sanitize_cell_text(r.get("seller")) or ("—" if r["status"] != checkpoint.STATUS_MISMATCHED else None),
        _sanitize_cell_text(r.get("product_title")),
        _STATUS_DISPLAY.get(r["status"], r["status"].upper()),
        _sanitize_cell_text(_issue_note(r)),
        "View on Amazon" if r.get("url") else None,
        _local_time(r.get("checked_at")),
    )


def _write_issue_sheet(ws: Worksheet, items: list[dict]) -> None:
    items = sorted(items, key=lambda r: (r["status"] != checkpoint.STATUS_MISMATCHED, r["asin"]))
    _write_table(
        ws, _ISSUE_HEADERS, [_issue_row(r) for r in items],
        widths={1: 14, 2: 12, 3: 12, 4: 12, 5: 11, 6: 26, 7: 50, 8: 16, 9: 46, 10: 15, 11: 21},
    )
    for i, r in enumerate(items, start=2):
        fill = MISMATCH_FILL if r["status"] == checkpoint.STATUS_MISMATCHED else LISTING_FILL
        for col in range(1, len(_ISSUE_HEADERS) + 1):
            ws.cell(row=i, column=col).fill = fill
        for col in _ISSUE_MONEY_COLS:
            ws.cell(row=i, column=col).number_format = MONEY_FORMAT
        if r.get("url"):
            link = ws.cell(row=i, column=_ISSUE_LINK_COL)
            link.hyperlink = r["url"]
            link.font = LINK_FONT


def _issue_items_for(items: list[dict], brand: str) -> list[dict]:
    return [r for r in items if _brand(r) == brand and r["status"] in checkpoint.ISSUE_STATUSES]


def _write_run_info(ws: Worksheet, run: dict, items: list[dict]) -> None:
    stats = {}
    if run.get("stats_json"):
        try:
            stats = json.loads(run["stats_json"])
        except (TypeError, ValueError):
            stats = {}
    by_resolver: dict[str, int] = {}
    for r in items:
        if r["status"] != checkpoint.STATUS_FAILED and r.get("resolved_by"):
            by_resolver[r["resolved_by"]] = by_resolver.get(r["resolved_by"], 0) + 1

    rows: list[tuple] = [
        ("Input file", run["input_filename"]),
        ("Started", _local_time(run["started_at"])),
        ("Finished", _local_time(run.get("finished_at"))),
        ("ASINs checked", len(items)),
        ("Price OK", run["matched"]),
        ("Price mismatch", run["mismatched"]),
        ("Listing issues (OOS / unavailable / no buy box / not found)", run["out_of_stock"]),
        ("Could not verify", run["failed"]),
        ("Tolerance", f"± ₹{run['tolerance_abs']:g}" + (f" or ± {run['tolerance_pct']:g}%" if run["tolerance_pct"] else "")),
    ]
    labels = {"http": "Resolved on first pass", "recovery": "Resolved on recovery pass",
              "offers": "Resolved via Amazon's offers page", "browser": "Resolved via Google Chrome check"}
    for key in ("http", "recovery", "offers", "browser"):
        if by_resolver.get(key):
            rows.append((labels[key], by_resolver[key]))
    oc = stats.get("offers_check") or {}
    if oc.get("allowed"):
        checked = oc.get("agreed", 0) + oc.get("disagreed", 0)
        if oc.get("enabled"):
            verdict = f"used — it agreed with the product page on {oc.get('agreed', 0)} of {checked} sample rows"
        elif oc.get("disagreed"):
            verdict = f"not used — it disagreed with the product page on {oc['disagreed']} of {checked} sample rows"
        else:
            verdict = "not used — not enough sample rows could be compared this run"
        rows.append(("Offers-page double-check", verdict))
    if stats.get("engine_restarts"):
        rows.append(("Recovered from internal errors", f"{stats['engine_restarts']} time(s) — no rows lost"))
    # Raw counters of the LATEST check only (a Resume / Retry / recovery
    # starts a new one), for support — labelled so they aren't read as totals.
    for key, value in stats.items():
        if key == "engine_restarts":
            continue
        if isinstance(value, (int, float, str)) and not isinstance(value, bool):
            rows.append((f"Diagnostics (latest check): {key.replace('_', ' ')}", value))
    _write_table(ws, ["Item", "Value"], [(_sanitize_cell_text(k), _sanitize_cell_text(v)) for k, v in rows],
                 widths={1: 52, 2: 40})


def build_report(run_id: str, output_path: Path | None = None, db_path: Path = config.DB_PATH) -> Path:
    run = checkpoint.get_run(run_id, db_path=db_path)
    if run is None:
        raise ValueError(f"Unknown run_id: {run_id}")

    items = checkpoint.get_run_items(run_id, db_path=db_path)
    brands = sorted({_brand(r) for r in items}, key=str.lower)
    failed = sorted((r for r in items if r["status"] == checkpoint.STATUS_FAILED), key=lambda r: r["asin"])

    wb = Workbook()

    ws_overview = wb.active
    ws_overview.title = "Overview"
    overview_rows = []
    for brand in brands:
        b_items = [r for r in items if _brand(r) == brand]
        matched = sum(1 for r in b_items if r["status"] == checkpoint.STATUS_MATCHED)
        mismatched = sum(1 for r in b_items if r["status"] == checkpoint.STATUS_MISMATCHED)
        listing = sum(1 for r in b_items if r["status"] in checkpoint.ISSUE_STATUSES
                      and r["status"] != checkpoint.STATUS_MISMATCHED)
        b_failed = sum(1 for r in b_items if r["status"] == checkpoint.STATUS_FAILED)
        overview_rows.append((_sanitize_cell_text(brand), len(b_items), matched, mismatched, listing, b_failed,
                              mismatched + listing))
    _write_table(
        ws_overview,
        ["Brand", "ASINs", "Price OK", "Price Mismatch", "Listing Issue (OOS / no buy box / not found)",
         "Could Not Verify", "Needs Action"],
        overview_rows,
        widths={1: 24, 2: 9, 3: 10, 4: 15, 5: 26, 6: 17, 7: 14},
    )

    used_sheet_names: set[str] = {"overview", "could not verify", "run info"}
    for brand in brands:
        ws = wb.create_sheet(_sanitize_sheet_name(brand, used_sheet_names))
        _write_issue_sheet(ws, _issue_items_for(items, brand))

    ws_failed = wb.create_sheet("Could Not Verify")
    _write_table(
        ws_failed,
        ["ASIN", "Brand", "Expected Price (₹)", "Reason", "Attempts", "Last Tried", "Amazon Link"],
        [(r["asin"], _sanitize_cell_text(_brand(r)), r["expected_price"],
          _sanitize_cell_text(_friendly_reason(r.get("error_reason")) or "Unknown"),
          r["attempts"], _local_time(r.get("checked_at")), "View on Amazon" if r.get("url") else None)
         for r in failed],
        widths={1: 14, 2: 20, 3: 12, 4: 60, 5: 10, 6: 21, 7: 15},
    )
    for i, r in enumerate(failed, start=2):
        ws_failed.cell(row=i, column=3).number_format = MONEY_FORMAT
        if r.get("url"):
            link = ws_failed.cell(row=i, column=7)
            link.hyperlink = r["url"]
            link.font = LINK_FONT
    summary_row = ws_failed.max_row + 2
    ws_failed.cell(row=summary_row, column=1, value=(
        f"{len(items) - len(failed)} of {len(items)} ASINs verified; {len(failed)} could not be verified."
        + (" Use Retry in the app to re-check them." if failed else "")
    )).font = Font(italic=True)

    _write_run_info(wb.create_sheet("Run Info"), run, items)

    if output_path is None:
        output_path = config.OUTPUT_DIR / f"PriceVerification_{run_id}.xlsx"
    return _save_workbook(wb, output_path)


def _save_workbook(wb: Workbook, output_path: Path) -> Path:
    """Save, or — if that file is open in Excel (Windows locks it; a Retry
    rebuilds the same run's report) — save under a timestamped name instead
    of failing."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(str(output_path))
        return output_path
    except PermissionError:
        alt = output_path.with_name(f"{output_path.stem}_{datetime.now():%Y%m%d_%H%M%S}{output_path.suffix}")
        wb.save(str(alt))
        return alt


def build_brand_report(run_id: str, brand: str, output_path: Path | None = None, db_path: Path = config.DB_PATH) -> Path:
    """One brand's issue list as its own small workbook — attach it directly
    to the email that goes to that brand's seller."""
    run = checkpoint.get_run(run_id, db_path=db_path)
    if run is None:
        raise ValueError(f"Unknown run_id: {run_id}")

    items = checkpoint.get_run_items(run_id, db_path=db_path)
    wb = Workbook()
    ws = wb.active
    ws.title = _sanitize_sheet_name(brand, set())
    _write_issue_sheet(ws, _issue_items_for(items, brand))

    if output_path is None:
        safe_brand = re.sub(r"[^A-Za-z0-9._-]+", "_", brand).strip("_") or "brand"
        output_path = config.OUTPUT_DIR / f"PriceIssues_{safe_brand}_{run_id}.xlsx"
    return _save_workbook(wb, output_path)
